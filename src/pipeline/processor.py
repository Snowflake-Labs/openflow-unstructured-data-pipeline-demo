"""
Document processing pipeline for extracting content from various file formats.
"""
import os
import logging
from typing import Dict, List, Optional, Any
from concurrent.futures import ThreadPoolExecutor, as_completed
import mimetypes

import PyPDF2
from docx import Document
import openpyxl
import pandas as pd

class DocumentProcessor:
    """Processor for extracting content from various document formats."""
    
    def __init__(self):
        self.logger = logging.getLogger(__name__)
        
    def extract_content(self, file_path: str, file_type: str = None) -> Dict[str, Any]:
        """
        Extract content from a document file.
        
        Args:
            file_path: Path to the document file
            file_type: MIME type of the file (optional)
            
        Returns:
            Dictionary containing extracted content and metadata
        """
        if not os.path.exists(file_path):
            self.logger.error(f"File not found: {file_path}")
            return {'error': 'File not found', 'content': ''}
        
        # Determine file type if not provided
        if not file_type:
            file_type, _ = mimetypes.guess_type(file_path)
        
        try:
            if file_type == 'application/pdf' or file_path.lower().endswith('.pdf'):
                return self._extract_pdf_content(file_path)
            elif file_type in ['application/vnd.openxmlformats-officedocument.wordprocessingml.document'] or file_path.lower().endswith('.docx'):
                return self._extract_docx_content(file_path)
            elif file_type in ['application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'] or file_path.lower().endswith('.xlsx'):
                return self._extract_xlsx_content(file_path)
            elif file_type in ['application/vnd.openxmlformats-officedocument.presentationml.presentation'] or file_path.lower().endswith('.pptx'):
                return self._extract_pptx_content(file_path)
            elif file_type == 'text/plain' or file_path.lower().endswith('.txt'):
                return self._extract_text_content(file_path)
            else:
                self.logger.warning(f"Unsupported file type: {file_type} for file: {file_path}")
                return {'error': f'Unsupported file type: {file_type}', 'content': ''}
                
        except Exception as e:
            self.logger.error(f"Error extracting content from {file_path}: {e}")
            return {'error': str(e), 'content': ''}
    
    def _extract_pdf_content(self, file_path: str) -> Dict[str, Any]:
        """Extract content from PDF files."""
        content_data = {
            'content': '',
            'pages': [],
            'metadata': {},
            'file_type': 'pdf'
        }
        
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                content_data['metadata'] = {
                    'num_pages': len(pdf_reader.pages),
                    'title': pdf_reader.metadata.get('/Title', '') if pdf_reader.metadata else '',
                    'author': pdf_reader.metadata.get('/Author', '') if pdf_reader.metadata else '',
                    'subject': pdf_reader.metadata.get('/Subject', '') if pdf_reader.metadata else ''
                }
                
                all_text = []
                for page_num, page in enumerate(pdf_reader.pages, 1):
                    try:
                        page_text = page.extract_text()
                        all_text.append(page_text)
                        content_data['pages'].append({
                            'page_number': page_num,
                            'content': page_text,
                            'word_count': len(page_text.split())
                        })
                    except Exception as e:
                        self.logger.warning(f"Error extracting page {page_num} from {file_path}: {e}")
                
                content_data['content'] = '\n\n'.join(all_text)
                content_data['word_count'] = len(content_data['content'].split())
                
        except Exception as e:
            content_data['error'] = str(e)
            
        return content_data
    
    def _extract_docx_content(self, file_path: str) -> Dict[str, Any]:
        """Extract content from DOCX files."""
        content_data = {
            'content': '',
            'paragraphs': [],
            'metadata': {},
            'file_type': 'docx'
        }
        
        try:
            doc = Document(file_path)
            
            # Extract core properties
            core_props = doc.core_properties
            content_data['metadata'] = {
                'title': core_props.title or '',
                'author': core_props.author or '',
                'subject': core_props.subject or '',
                'created': str(core_props.created) if core_props.created else '',
                'modified': str(core_props.modified) if core_props.modified else '',
                'num_paragraphs': len(doc.paragraphs)
            }
            
            # Extract paragraph content
            all_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    all_text.append(para.text)
                    content_data['paragraphs'].append({
                        'text': para.text,
                        'style': para.style.name if para.style else 'Normal'
                    })
            
            content_data['content'] = '\n\n'.join(all_text)
            content_data['word_count'] = len(content_data['content'].split())
            
        except Exception as e:
            content_data['error'] = str(e)
            
        return content_data
    
    def _extract_xlsx_content(self, file_path: str) -> Dict[str, Any]:
        """Extract content from Excel files."""
        content_data = {
            'content': '',
            'sheets': [],
            'metadata': {},
            'file_type': 'xlsx'
        }
        
        try:
            workbook = openpyxl.load_workbook(file_path, data_only=True)
            
            content_data['metadata'] = {
                'num_sheets': len(workbook.worksheets),
                'sheet_names': workbook.sheetnames
            }
            
            all_content = []
            
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                sheet_data = {
                    'name': sheet_name,
                    'content': '',
                    'rows': sheet.max_row,
                    'columns': sheet.max_column
                }
                
                # Extract data from sheet
                sheet_content = []
                for row in sheet.iter_rows(values_only=True):
                    row_text = '\t'.join([str(cell) if cell is not None else '' for cell in row])
                    if row_text.strip():
                        sheet_content.append(row_text)
                
                sheet_data['content'] = '\n'.join(sheet_content)
                content_data['sheets'].append(sheet_data)
                all_content.append(f"Sheet: {sheet_name}\n{sheet_data['content']}")
            
            content_data['content'] = '\n\n'.join(all_content)
            content_data['word_count'] = len(content_data['content'].split())
            
        except Exception as e:
            content_data['error'] = str(e)
            
        return content_data
    
    def _extract_pptx_content(self, file_path: str) -> Dict[str, Any]:
        """Extract content from PowerPoint files."""
        content_data = {
            'content': '',
            'slides': [],
            'metadata': {},
            'file_type': 'pptx'
        }
        
        try:
            from pptx import Presentation
            
            prs = Presentation(file_path)
            
            content_data['metadata'] = {
                'num_slides': len(prs.slides),
                'title': prs.core_properties.title or '',
                'author': prs.core_properties.author or '',
                'subject': prs.core_properties.subject or ''
            }
            
            all_text = []
            
            for slide_num, slide in enumerate(prs.slides, 1):
                slide_text = []
                
                for shape in slide.shapes:
                    if hasattr(shape, 'text') and shape.text.strip():
                        slide_text.append(shape.text)
                
                slide_content = '\n'.join(slide_text)
                content_data['slides'].append({
                    'slide_number': slide_num,
                    'content': slide_content,
                    'word_count': len(slide_content.split())
                })
                
                if slide_content:
                    all_text.append(f"Slide {slide_num}:\n{slide_content}")
            
            content_data['content'] = '\n\n'.join(all_text)
            content_data['word_count'] = len(content_data['content'].split())
            
        except ImportError:
            content_data['error'] = 'python-pptx library not available for PowerPoint processing'
        except Exception as e:
            content_data['error'] = str(e)
            
        return content_data
    
    def _extract_text_content(self, file_path: str) -> Dict[str, Any]:
        """Extract content from plain text files."""
        content_data = {
            'content': '',
            'metadata': {},
            'file_type': 'text'
        }
        
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as file:
                content = file.read()
                
            content_data['content'] = content
            content_data['metadata'] = {
                'line_count': len(content.split('\n')),
                'word_count': len(content.split()),
                'char_count': len(content)
            }
            content_data['word_count'] = content_data['metadata']['word_count']
            
        except Exception as e:
            content_data['error'] = str(e)
            
        return content_data
    
    def chunk_content(self, content: str, chunk_size: int = 1000, overlap: int = 100) -> List[Dict[str, Any]]:
        """
        Split content into chunks for processing.
        
        Args:
            content: Text content to chunk
            chunk_size: Maximum size of each chunk
            overlap: Number of characters to overlap between chunks
            
        Returns:
            List of content chunks with metadata
        """
        if not content or len(content) <= chunk_size:
            return [{
                'chunk_number': 1,
                'content': content,
                'start_pos': 0,
                'end_pos': len(content),
                'size': len(content)
            }]
        
        chunks = []
        start = 0
        chunk_num = 1
        
        while start < len(content):
            end = min(start + chunk_size, len(content))
            
            # Try to break at sentence or word boundary
            if end < len(content):
                # Look for sentence boundary
                last_period = content.rfind('.', start, end)
                last_newline = content.rfind('\n', start, end)
                last_space = content.rfind(' ', start, end)
                
                boundary = max(last_period, last_newline, last_space)
                if boundary > start + chunk_size // 2:  # Don't make chunks too small
                    end = boundary + 1
            
            chunk_content = content[start:end].strip()
            if chunk_content:
                chunks.append({
                    'chunk_number': chunk_num,
                    'content': chunk_content,
                    'start_pos': start,
                    'end_pos': end,
                    'size': len(chunk_content)
                })
                chunk_num += 1
            
            start = max(start + 1, end - overlap)
        
        return chunks
    
    def process_documents_batch(self, file_paths: List[str], max_workers: int = 4) -> List[Dict[str, Any]]:
        """
        Process multiple documents in parallel.
        
        Args:
            file_paths: List of file paths to process
            max_workers: Maximum number of worker threads
            
        Returns:
            List of processing results
        """
        results = []
        
        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            future_to_path = {
                executor.submit(self.extract_content, path): path 
                for path in file_paths
            }
            
            for future in as_completed(future_to_path):
                path = future_to_path[future]
                try:
                    result = future.result()
                    result['file_path'] = path
                    results.append(result)
                    self.logger.info(f"Processed: {os.path.basename(path)}")
                except Exception as e:
                    self.logger.error(f"Error processing {path}: {e}")
                    results.append({
                        'file_path': path,
                        'error': str(e),
                        'content': ''
                    })
        
        return results
    
    def get_content_summary(self, content: str, max_length: int = 500) -> str:
        """
        Generate a summary of the content for quick review.
        
        Args:
            content: Full content text
            max_length: Maximum length of summary
            
        Returns:
            Content summary
        """
        if not content:
            return "No content available"
        
        # Simple extractive summary - get first few sentences up to max_length
        sentences = content.replace('\n', ' ').split('.')
        summary = ""
        
        for sentence in sentences:
            sentence = sentence.strip()
            if sentence and len(summary + sentence) < max_length:
                summary += sentence + ". "
            else:
                break
        
        if not summary and content:
            # Fallback to first max_length characters
            summary = content[:max_length].strip() + "..."
        
        return summary.strip()
    
    def identify_key_sections(self, content: str, file_type: str = 'text') -> Dict[str, str]:
        """
        Identify key business document sections.
        
        Args:
            content: Document content
            file_type: Type of document
            
        Returns:
            Dictionary of identified sections
        """
        sections = {}
        content_lower = content.lower()
        
        # Common business document section patterns
        section_patterns = {
            'executive_summary': ['executive summary', 'summary', 'overview'],
            'objectives': ['objectives', 'goals', 'targets'],
            'strategy': ['strategy', 'strategic', 'approach'],
            'financial': ['financial', 'budget', 'revenue', 'cost', 'profit'],
            'risks': ['risk', 'threat', 'challenge'],
            'recommendations': ['recommendation', 'action item', 'next step'],
            'conclusion': ['conclusion', 'summary', 'closing']
        }
        
        for section_name, keywords in section_patterns.items():
            for keyword in keywords:
                if keyword in content_lower:
                    # Extract surrounding context
                    start_pos = content_lower.find(keyword)
                    if start_pos != -1:
                        # Get paragraph containing the keyword
                        para_start = content.rfind('\n', 0, start_pos)
                        para_end = content.find('\n\n', start_pos)
                        if para_end == -1:
                            para_end = start_pos + 500
                        
                        section_content = content[para_start:para_end].strip()
                        if len(section_content) > 50:  # Minimum content length
                            sections[section_name] = section_content[:500]  # Limit length
                            break
        
        return sections